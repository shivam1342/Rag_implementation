# app/services/document_service.py
import os
from typing import List
from app.core.config import settings
from app.core.logger import logger

class DocumentService:
    """Service for document processing and text extraction"""
    
    def __init__(self):
        os.makedirs(settings.UPLOAD_DIR, exist_ok=True)
        logger.info("Document service initialized")
    
    def extract_text(self, file_path: str) -> str:
        """
        Extract text from uploaded file.
        Supports: TXT, PDF, DOCX
        
        Args:
            file_path: Path to the file
            
        Returns:
            Extracted text
        """
        try:
            file_ext = os.path.splitext(file_path)[1].lower()
            
            if file_ext == ".txt":
                return self._extract_from_txt(file_path)
            elif file_ext == ".pdf":
                return self._extract_from_pdf(file_path)
            elif file_ext == ".docx":
                return self._extract_from_docx(file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_ext}")
        except Exception as e:
            logger.error(f"Error extracting text from {file_path}: {e}")
            raise
    
    def _extract_from_txt(self, file_path: str) -> str:
        """Extract text from TXT file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()
        logger.info(f"Extracted {len(text)} characters from TXT file")
        return text
    
    def _extract_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file using pypdf"""
        try:
            from pypdf import PdfReader
            
            reader = PdfReader(file_path)
            text = ""
            
            for page_num, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            
            logger.info(f"Extracted {len(text)} characters from PDF ({len(reader.pages)} pages)")
            return text
        except ImportError:
            raise ImportError("pypdf library not installed. Run: pip install pypdf")
        except Exception as e:
            logger.error(f"Error extracting PDF: {e}")
            raise
    
    def _extract_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX file using python-docx"""
        try:
            from docx import Document
            
            doc = Document(file_path)
            text = ""
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + "\t"
                    text += "\n"
            
            logger.info(f"Extracted {len(text)} characters from DOCX ({len(doc.paragraphs)} paragraphs)")
            return text
        except ImportError:
            raise ImportError("python-docx library not installed. Run: pip install python-docx")
        except Exception as e:
            logger.error(f"Error extracting DOCX: {e}")
            raise
    
    def save_upload(self, filename: str, content: bytes) -> str:
        """
        Save uploaded file to disk.
        
        Args:
            filename: Name of the file
            content: File content as bytes
            
        Returns:
            Path to saved file
        """
        file_path = os.path.join(settings.UPLOAD_DIR, filename)
        
        with open(file_path, 'wb') as f:
            f.write(content)
        
        logger.info(f"Saved file: {file_path}")
        return file_path
    
    def list_uploaded_files(self) -> List[str]:
        """List all uploaded files"""
        try:
            files = os.listdir(settings.UPLOAD_DIR)
            return [f for f in files if os.path.isfile(os.path.join(settings.UPLOAD_DIR, f))]
        except Exception as e:
            logger.error(f"Error listing files: {e}")
            return []

# Global instance
document_service = DocumentService()
